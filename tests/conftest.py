"""
pytest 会话级配置
=================
每次测试会话结束后自动清理运行时缓存（保留模板文件本身）。
"""

import os
import shutil
import pytest


@pytest.fixture(scope="session", autouse=True)
def auto_clean_cache(request):
    """
    整个测试会话结束后自动清理运行时生成的缓存文件，
    避免 data/ 目录堆积日志、任务状态和附件等测试产物。

    行为：
    - logs/ tasks/ attachments/ → 清空目录
    - *.xlsx → 仅清除数据行，保留模板结构和表头
    """

    def _clear_excel_data(filepath: str) -> None:
        """保留 Excel 模板的前 N 行（标题+表头），清除后面的数据行"""
        try:
            import openpyxl
            wb = openpyxl.load_workbook(filepath)
            for ws in wb.worksheets:
                max_r = ws.max_row or 0
                max_c = ws.max_column or 0
                if max_r == 0:
                    continue

                # 从头扫描每一行，找到包含表头关键词的行即停止
                # 注意：使用单元格级别匹配而非拼接文本匹配，避免描述性文字中
                # 意外包含关键词（如"系统将自动生成序号"中的"序号"）导致误判。
                header_row = 1  # 兜底：至少保留第1行
                header_keywords = ("姓名*", "立项通过日期", "荣誉时数值*", "活动名称", "序号")
                for r in range(1, max_r + 1):
                    matched_cells = 0
                    for c in range(1, max_c + 1):
                        val = ws.cell(row=r, column=c).value
                        if val is not None:
                            val_str = str(val).strip()
                            # 表头单元格通常是简短的关键词（≤20 字符），
                            # 而说明性文字较长，排除掉以避免误匹配
                            if len(val_str) <= 20:
                                if any(kw in val_str for kw in header_keywords):
                                    matched_cells += 1
                    # 至少匹配到 2 个表头关键词才认为是表头行
                    if matched_cells >= 2:
                        header_row = r
                        break

                # 删除表头行之后的所有数据行（真正删除，而非设 None）
                if max_r > header_row:
                    ws.delete_rows(header_row + 1, max_r - header_row)

            wb.save(filepath)
        except Exception as e:
            import traceback
            print(f"  [clean_cache] 清理 Excel 数据失败 {filepath}: {e}")
            traceback.print_exc()

    def cleanup():
        import logging as _logging

        project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        data_dir = os.path.join(project_root, "data")

        # ── 0. 先释放所有 logging FileHandler 持有的文件句柄 ──
        #     Windows 上打开的文件无法被 os.unlink/shtuil.rmtree 删除，
        #     必须在清理目录之前关闭 handler。
        root_logger = _logging.getLogger()
        for handler in root_logger.handlers[:]:
            if isinstance(handler, _logging.FileHandler):
                handler.close()
                root_logger.removeHandler(handler)

        # ── 1. 清理 Python 字节码缓存 ──────────────────
        import glob as _glob

        # 删除所有 __pycache__ 目录
        for pycache in _glob.glob(os.path.join(project_root, "**", "__pycache__"), recursive=True):
            try:
                shutil.rmtree(pycache)
            except Exception as e:
                print(f"  [clean_cache] 清理失败 {pycache}: {e}")

        # 删除所有 .pyc/.pyo 编译文件
        for pattern in ("**/*.pyc", "**/*.pyo"):
            for fpath in _glob.glob(os.path.join(project_root, pattern), recursive=True):
                try:
                    os.unlink(fpath)
                except Exception as e:
                    print(f"  [clean_cache] 清理失败 {fpath}: {e}")

        # 删除 .pytest_cache 目录
        pytest_cache = os.path.join(project_root, ".pytest_cache")
        if os.path.isdir(pytest_cache):
            try:
                shutil.rmtree(pytest_cache)
            except Exception as e:
                print(f"  [clean_cache] 清理失败 {pytest_cache}: {e}")

        # ── 2. 清理运行时目录（日志、任务状态、附件缓存等）──
        runtime_dirs = [
            os.path.join(data_dir, "logs"),
            os.path.join(data_dir, "tasks"),
            os.path.join(data_dir, "attachments"),
        ]

        for dir_path in runtime_dirs:
            if os.path.isdir(dir_path):
                for fname in os.listdir(dir_path):
                    fpath = os.path.join(dir_path, fname)
                    try:
                        if os.path.isfile(fpath) or os.path.islink(fpath):
                            os.unlink(fpath)
                        elif os.path.isdir(fpath):
                            shutil.rmtree(fpath)
                    except Exception as e:
                        print(f"  [clean_cache] 清理失败 {fpath}: {e}")

        # ── 3. 清理 Excel 模板中的数据行（保留模板文件本身）──
        for fpath in _glob.glob(os.path.join(data_dir, "*.xlsx")):
            _clear_excel_data(fpath)

    request.addfinalizer(cleanup)
